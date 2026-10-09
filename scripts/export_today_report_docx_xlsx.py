# -*- coding: utf-8 -*-
"""生成今日采集/上架情况 Word + Excel（给 Boss）。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter

OUT_DIR = Path(r"E:\tk-miaoshou-rework\clean_rebuild\data\previews")
STAMP = datetime.now().strftime("%Y%m%d_%H%M")
WORD_PATH = OUT_DIR / f"今日采集上架情况说明_{STAMP}.docx"
XLSX_PATH = OUT_DIR / f"今日采集上架数据_{STAMP}.xlsx"


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 13, bold=True)
    return p


def add_para(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def build_word():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TikTok MX 店铺采集与上架日报")
    set_run_font(run, size=18, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"日期：2026-09-17（墨西哥站）｜生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    set_run_font(run, size=10, color=(90, 90, 90))

    add_heading_cn(doc, "一、核心结论", 1)
    add_para(
        doc,
        "今天两店合计成功上架 155 条（小赵 96 + 小韩 59）。该数字已与妙手「发布记录」逐条核对一致，"
        "不是脚本自行统计的虚数。日限为单店 300，远未触顶。上架偏少的主因是：飞书链接经妙手 OpenAPI 采集成功率偏低"
        "（可统计批量约 49%），采进后再被门禁留库与改品并发冲突进一步损耗。",
    )
    add_para(
        doc,
        "补充核实：老板/运营若中间有手动操作，以妙手发布记录为准——今天账号内成功发布恰好 155 条，"
        "未发现脚本之外的额外成功发布增量。飞书回写「上架=是」少于 155，是因为部分公共箱直发未回写飞书，"
        "不能代表真实发布量。",
        bold=False,
    )

    add_heading_cn(doc, "二、妙手核实 vs 脚本计数", 1)
    add_table(
        doc,
        ["口径", "小赵1店", "小韩1店", "合计", "说明"],
        [
            ["妙手发布记录（今天 success）", "96", "59", "155", "最权威，已按 shopId 拆分"],
            ["本地自动化脚本配额", "96", "59", "155", "与妙手完全一致"],
            ["飞书「上架=是」且今天有改动", "96", "40", "136", "韩少计约 19（未回写）"],
            ["单店日限", "300", "300", "600", "未用满"],
        ],
    )
    add_para(doc, "小赵发布时段分布：17 点 40 条、18 点 48 条、19 点 8 条。", size=10)
    add_para(doc, "小韩发布时段分布：19 点 13 条、20 点 45 条、21 点 1 条。", size=10)

    add_heading_cn(doc, "三、为什么上架少（因果链）", 1)
    add_para(doc, "1）采集层（最大损耗）", bold=True)
    add_para(
        doc,
        "流程：飞书短链解析 → 调用妙手 fetch_item → 公共箱出现 success 才算采成。"
        "今天大量链接在公共箱落成 status=fail，妙手返回原因统一为："
        "「建议采集1688产品或使用插件采集」。含义是妙手服务器端拉取 TikTok MX 商品失败，"
        "并非脚本未提交。飞书表后半段更差：小韩两批 API 采集率仅约 32%～36%。",
    )
    add_para(doc, "2）门禁留库", bold=True)
    add_para(
        doc,
        "采成功后仍可能因规则不上架：类目/受众不符（童装、帽衫等）、标题与规格无具体颜色"
        "（仅允许黑/白/红/蓝/粉）、多件装等。例如公共箱→小韩一批 26 条中留库 10 条（约 38%）。",
    )
    add_para(doc, "3）改品并发冲突", bold=True)
    add_para(
        doc,
        "保存改品时妙手报错「编辑过程中产品数据发生变动…请重新打开弹窗」。"
        "小赵批量 100 中上架失败 19 条，多数为此类；小韩批次也有发生。属于平台侧并发/版本冲突。",
    )
    add_para(doc, "简化漏斗：", bold=True)
    add_para(
        doc,
        "飞书可统计取链 272 条 → 采集成功约 134 条（49%）→ 叠加 until_ok / 公共箱批次后两店最终发布 155 条。"
        "日限 600 条额度几乎没用满，瓶颈在「采得进 + 改得过」，不在配额。",
    )

    add_heading_cn(doc, "四、分批执行明细（摘要）", 1)
    add_table(
        doc,
        ["批次", "店铺", "入口", "采集成功", "采集率", "上架成功", "留库", "失败", "端到端"],
        [
            ["until_ok", "小赵", "尝试 105", "—", "—", "44", "6", "55", "42%"],
            ["飞书批量 100", "小赵", "100", "75", "75%", "52", "4", "19", "52%"],
            ["公共箱→店（含手动采）", "小韩", "池 26", "—", "—", "13", "10", "3", "50%"],
            ["飞书 1066–1138", "小韩", "72", "23", "32%", "22", "9", "2", "31%"],
            ["飞书 1139–1238", "小韩", "100", "36", "36%", "23", "2", "11", "23%"],
            ["晚间补采 recent", "小韩", "池 2", "—", "—", "1", "0", "1", "50%"],
        ],
    )

    add_heading_cn(doc, "五、可统计采集合计（飞书 API 三批）", 1)
    add_table(
        doc,
        ["指标", "数量", "比例"],
        [
            ["取链合计", "272", "100%"],
            ["采集成功", "134", "49%"],
            ["采集失败", "138", "51%"],
            ["其中小赵批量采集率", "75/100", "75%"],
            ["其中小韩两批采集率", "59/172", "34%"],
        ],
    )
    add_para(
        doc,
        "采集失败典型原因（抽检公共箱 fail 记录）：全部为妙手提示改走 1688 或浏览器插件采集。"
        "插件采集走浏览器登录态，成功率通常高于 OpenAPI，但无法无人值守；今日晚间补采仍出现大量 fail。",
        size=10,
    )

    add_heading_cn(doc, "六、建议与后续", 1)
    add_para(doc, "1. 飞书后半段链接质量需筛选，或对 API 采失败链接改用插件补采后再走认领上架。")
    add_para(doc, "2. 改品并发失败可做有限次自动重试，减少「采到了却上不去」。")
    add_para(doc, "3. 对外汇报以上架数请以「妙手发布记录」为准；飞书回写仅作过程跟踪。")
    add_para(doc, "4. 详细分项数字见同目录 Excel《今日采集上架数据》。")

    add_para(doc, "— 文档结束 —", size=10)
    doc.save(WORD_PATH)
    return WORD_PATH


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def style_body(ws, start=2):
    font = Font(name="微软雅黑", size=10)
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    for row in ws.iter_rows(min_row=start, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = font
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_excel():
    wb = Workbook()

    # Sheet1 总览
    ws = wb.active
    ws.title = "总览核实"
    ws.append(["项目", "小赵1店", "小韩1店", "合计", "备注"])
    ws.append(["妙手发布记录今天 success", 96, 59, 155, "按 shopId 拆分；账号今日全部成功发布"])
    ws.append(["本地脚本配额", 96, 59, 155, "与妙手一致"])
    ws.append(["飞书上架=是且今天改动", 96, 40, 136, "韩少计因部分未回写飞书"])
    ws.append(["单店日限", 300, 300, 600, "未触顶"])
    ws.append(["脚本端到端上架合计", 96, 59, 155, "44+52；13+22+23+1"])
    style_header(ws)
    style_body(ws)
    autosize(ws, [28, 14, 14, 10, 42])
    ws.row_dimensions[1].height = 22

    # Sheet2 分批
    ws2 = wb.create_sheet("分批明细")
    ws2.append(
        [
            "批次",
            "店铺",
            "入口数量",
            "采集成功",
            "采集失败",
            "采集率",
            "上架池/处理数",
            "上架成功",
            "留库",
            "失败",
            "端到端上架率",
            "备注",
        ]
    )
    batches = [
        ["until_ok", "小赵1店", 105, "",  "", "", 105, 44, 6, 55, "42%", "逐条采+上；失败含采集"],
        ["飞书批量100", "小赵1店", 100, 75, 25, "75%", 75, 52, 4, 19, "52%", "采成后上架率约69%"],
        ["公共箱→韩", "小韩1店", 26, "", "", "", 26, 13, 10, 3, "50%", "含手动采集后的 success"],
        ["飞书1066-1138", "小韩1店", 72, 23, 49, "32%", 33, 22, 9, 2, "31%", "池含公共箱额外合并"],
        ["飞书1139-1238", "小韩1店", 100, 36, 64, "36%", 36, 23, 2, 11, "23%", ""],
        ["补采recent", "小韩1店", 2, "", "", "", 2, 1, 0, 1, "50%", "21:10后公共箱 success"],
    ]
    for r in batches:
        ws2.append(r)
    ws2.append([])
    ws2.append(["合计（脚本上架）", "", "", "", "", "", "", 155, "", "", "", "赵96+韩59"])
    ws2.append(["飞书API三批取链", "赵批+韩两批", 272, 134, 138, "49%", "", "", "", "", "", "可统计采集漏斗"])
    style_header(ws2)
    style_body(ws2)
    autosize(ws2, [16, 10, 10, 10, 10, 10, 12, 10, 8, 8, 12, 28])

    # Sheet3 采集漏斗
    ws3 = wb.create_sheet("采集漏斗")
    ws3.append(["环节", "数量", "占上一环/占比", "说明"])
    ws3.append(["飞书API可统计取链", 272, "100%", "赵批量100 + 韩72 + 韩100"])
    ws3.append(["采集成功", 134, "49%", "公共箱 status=success"])
    ws3.append(["采集失败", 138, "51%", "公共箱 fail；妙手建议插件/1688"])
    ws3.append(["小赵批量采集成功", 75, "75% of 100", "相对较好"])
    ws3.append(["小韩两批采集成功", 59, "34% of 172", "飞书后段更差"])
    ws3.append(["采成后上架（约）", "", "约65%-70%", "扣门禁+改品并发"])
    ws3.append(["两店最终发布（妙手）", 155, "—", "含 until_ok 与公共箱批次"])
    style_header(ws3)
    style_body(ws3)
    autosize(ws3, [22, 14, 16, 40])

    # Sheet4 失败原因
    ws4 = wb.create_sheet("失败与留库原因")
    ws4.append(["类型", "典型表现", "主要出现批次", "对上架的影响"])
    ws4.append(
        [
            "采集失败",
            "公共箱 fail；文案：建议采集1688或使用插件",
            "韩飞书两批、赵until_ok、晚间补采",
            "最大头；链接解析成功也不等于采成",
        ]
    )
    ws4.append(
        [
            "改品并发",
            "编辑过程中产品数据发生变动，请重新打开弹窗",
            "赵批量（上架失败19中多数）、韩多批",
            "采到了仍上不去",
        ]
    )
    ws4.append(
        [
            "门禁-类目/受众",
            "wrong_category_or_audience",
            "各批留库",
            "童装/帽衫/类目不符等",
        ]
    )
    ws4.append(
        [
            "门禁-颜色",
            "title_and_options_no_concrete_color",
            "各批留库",
            "颜色不在黑白红蓝粉",
        ]
    )
    ws4.append(["SKU异常", "SKU不存在 / skuMap数据异常", "赵批量少量", "占比较小"])
    ws4.append(["短链SSL", "解析 shop.tiktok.com SSL/超时", "赵until_ok少量", "占比较小"])
    style_header(ws4)
    style_body(ws4)
    autosize(ws4, [16, 42, 28, 28])

    # Sheet5 时段
    ws5 = wb.create_sheet("发布时段")
    ws5.append(["店铺", "小时(本地+8)", "成功发布条数"])
    for h, n in [("17", 40), ("18", 48), ("19", 8)]:
        ws5.append(["小赵1店", h, n])
    for h, n in [("19", 13), ("20", 45), ("21", 1)]:
        ws5.append(["小韩1店", h, n])
    ws5.append(["小赵合计", "", 96])
    ws5.append(["小韩合计", "", 59])
    style_header(ws5)
    style_body(ws5)
    autosize(ws5, [12, 14, 14])

    # Sheet6 说明
    ws6 = wb.create_sheet("字段说明")
    ws6.append(["字段/口径", "含义"])
    ws6.append(["妙手发布记录", "search_move_collect_list，status=success，按 gmtCreate 归属今天"])
    ws6.append(["本地脚本配额", "data/daily_quota.json，脚本每成功上架+1"])
    ws6.append(["采集成功", "公共箱出现对应货源的 success 记录"])
    ws6.append(["留库", "门禁拦截，改品可不保存，不上架"])
    ws6.append(["端到端上架率", "最终上架成功 / 入口取链或尝试数"])
    ws6.append(["日限", "配置 daily_limit_per_shop=300 / 店"])
    style_header(ws6)
    style_body(ws6)
    autosize(ws6, [18, 70])

    wb.save(XLSX_PATH)
    return XLSX_PATH


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    w = build_word()
    x = build_excel()
    print(w)
    print(x)


if __name__ == "__main__":
    main()
