# -*- coding: utf-8 -*-
"""卫衣专线：公共箱 success → 文字+识图筛可用（与半袖 T 线完全独立）。

规则摘要（scheme=hoodie_mx）：
- 必须是卫衣（sudadera/hoodie/sweatshirt/卫衣/连帽）
- 底色只留黑/白（纯黑；水洗黑/灰黑删）
- 袖子须与衣身同色纯色，袖面不能有额外印花/撞色
- 只删不补，不填 Negro/Blanco
"""
from __future__ import annotations

import json
import shutil
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from api import endpoints as ep
from api.miaoshou_client import MiaoshouClient
from rules.product_gates import HOODIE_MARKERS
from services.auto_batch import qps_retry
from services.transformer import transform_product

TZ = timezone(timedelta(hours=8))
WORKERS = 4
SCHEME_PATH = ROOT / "config/schemes/hoodie_mx.json"


def parse_gmt(s: str) -> datetime | None:
    s = str(s or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.strptime(s[:26], fmt).replace(tzinfo=TZ)
        except ValueError:
            continue
    return None


def adapt_public_product(prod: dict) -> dict:
    out = dict(prod)
    color_name = str(out.get("colorPropName") or "Color")
    size_name = str(out.get("sizePropName") or "Tamaño")
    color_vals = []
    for key, val in (out.get("colorMap") or {}).items():
        name = val.get("name") if isinstance(val, dict) else key
        color_vals.append({"attrValue": str(name or key), "attrValueId": str(key)})
    size_vals = []
    for key, val in (out.get("sizeMap") or {}).items():
        name = val.get("name") if isinstance(val, dict) else key
        size_vals.append({"attrValue": str(name or key), "attrValueId": str(key)})
    props = []
    if color_vals:
        props.append({"name": color_name, "attrValueList": color_vals})
    if size_vals:
        props.append({"name": size_name, "attrValueList": size_vals})
    out["skuPropertyList"] = props
    return out


def title_looks_hoodie(title: str) -> bool:
    n = (title or "").strip().lower()
    if not n:
        return False
    return any(m in n for m in HOODIE_MARKERS)


def list_success_since(client: MiaoshouClient, since: datetime) -> list[dict]:
    rows: dict[int, dict] = {}
    for page in range(1, 120):
        body = qps_retry(
            lambda p=page: client.assert_success(
                client.post(ep.PUBLIC_LIST, {"pageNo": p, "pageSize": 100}),
                "公共箱",
            ),
            "public_list",
        )
        batch = (body.get("data") or {}).get("detailList") or []
        if not batch:
            break
        for it in batch:
            if str(it.get("status") or "").lower() != "success":
                continue
            gmt = parse_gmt(str(it.get("gmtCreate") or ""))
            if not gmt or gmt < since:
                continue
            cid = it.get("commonCollectBoxDetailId")
            sid = str(it.get("itemNum") or "")
            if not sid:
                for s in it.get("sourceList") or []:
                    if isinstance(s, dict) and s.get("sourceItemId"):
                        sid = str(s["sourceItemId"])
                        break
            if not sid or cid is None:
                continue
            rows[int(cid)] = {
                "common_id": int(cid),
                "sid": sid,
                "title": str(it.get("title") or "")[:160],
                "gmt": str(it.get("gmtCreate") or ""),
            }
        oldest = parse_gmt(str(batch[-1].get("gmtCreate") or ""))
        if oldest and oldest < since:
            break
        if len(batch) < 100:
            break
        time.sleep(0.12)
    return sorted(rows.values(), key=lambda x: x.get("gmt") or "", reverse=True)


def screen_one(client, row, scheme, blank, fill_urls, vision_cfg, detail_lock) -> dict:
    sid = str(row["sid"])
    out = {
        "source_id": sid,
        "common_id": row["common_id"],
        "gmt": row.get("gmt"),
        "title": row.get("title") or "",
        "usable": False,
        "stage": "",
        "reason": "",
        "images_scanned": 0,
        "prompt_tokens": 0,
        "completion_tokens": 0,
        "total_tokens": 0,
        "cny": 0.0,
        "cache_hits": 0,
        "error": "",
        "line": "hoodie",
    }
    try:
        with detail_lock:
            body = qps_retry(
                lambda cid=row["common_id"]: client.post(
                    ep.PUBLIC_DETAIL, {"commonCollectBoxDetailId": int(cid)}
                ),
                "detail",
            )
            time.sleep(0.25)
        inner = (body.get("body") or body).get("data") or {}
        if "result" in (body.get("body") or {}) and (body.get("body") or {}).get("result") != "success":
            out["stage"] = "detail"
            out["reason"] = str((body.get("body") or {}).get("message") or "detail_fail")[:160]
            return out
        prod = inner.get("editCommonCollectBoxDetail") or {}
        if not prod:
            out["stage"] = "detail"
            out["reason"] = "no_detail"
            return out
        title = str(prod.get("oriTitle") or prod.get("title") or row.get("title") or "")
        out["title"] = title[:160]

        full = transform_product(
            adapt_public_product(prod),
            scheme,
            blank,
            fill_image_urls=fill_urls,
            vision_cfg=vision_cfg,
        )
        hold = full.get("hold_without_publish") or {}
        vh = hold.get("vision") if isinstance(hold, dict) else {}
        if not isinstance(vh, dict):
            vh = {}
        out["images_scanned"] = int(vh.get("images_scanned") or 0)
        out["prompt_tokens"] = int(vh.get("prompt_tokens") or 0)
        out["completion_tokens"] = int(vh.get("completion_tokens") or 0)
        out["total_tokens"] = int(
            vh.get("total_tokens") or (out["prompt_tokens"] + out["completion_tokens"])
        )
        out["cny"] = float(vh.get("cny") or 0)
        out["cache_hits"] = int(vh.get("cache_hits") or 0)

        if hold.get("hold"):
            reason = str(hold.get("reason") or "hold")
            skipped = str(vh.get("reason") or "")
            if "skipped_pending_text" in skipped or (
                out["images_scanned"] == 0 and skipped.startswith("skipped")
            ):
                out["stage"] = "text"
            elif out["images_scanned"] or str(vh.get("gate") or "") == "vision_delete_only":
                out["stage"] = "vision"
            else:
                out["stage"] = "hold"
            out["reason"] = reason[:200]
            return out

        out["usable"] = True
        out["stage"] = "ok"
        out["reason"] = "ok"
        return out
    except Exception as e:
        out["stage"] = "error"
        out["error"] = str(e)[:240]
        out["reason"] = f"error:{type(e).__name__}"
        return out


def write_report(results: list[dict], pool_n: int, stamp: str, out_dir: Path, model: str) -> Path:
    usable = [r for r in results if r.get("usable")]
    text_hold = [r for r in results if r.get("stage") == "text"]
    vision_hold = [r for r in results if r.get("stage") == "vision"]
    other_hold = [r for r in results if r.get("stage") in ("hold", "detail") and not r.get("usable")]
    errors = [r for r in results if r.get("stage") == "error"]

    pt = sum(int(r.get("prompt_tokens") or 0) for r in results)
    ct = sum(int(r.get("completion_tokens") or 0) for r in results)
    tt = sum(int(r.get("total_tokens") or 0) for r in results) or (pt + ct)
    imgs = sum(int(r.get("images_scanned") or 0) for r in results)
    cny = sum(float(r.get("cny") or 0) for r in results)
    cache = sum(int(r.get("cache_hits") or 0) for r in results)
    visioned = sum(
        1 for r in results if int(r.get("images_scanned") or 0) > 0 or r.get("stage") in ("vision", "ok")
    )
    reasons = Counter(str(r.get("reason") or "")[:80] for r in results if not r.get("usable"))

    desk = Path.home() / "Desktop"
    name = f"卫衣识图筛可用报告_{stamp[:8]}.xlsx"
    path_prev = out_dir / name
    path_desk = desk / name

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    label_fill = PatternFill("solid", fgColor="D6EAF8")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    summary = [
        ["指标", "数值", "说明"],
        ["专线", "卫衣 hoodie_mx", "半袖 T 线未改动"],
        ["日期", stamp[:8], ""],
        ["候选池（标题含卫衣）", pool_n, "公共箱 success 预筛"],
        ["已筛条数", len(results), ""],
        ["可用（过文字+识图）", len(usable), "黑白底+同色袖"],
        ["文字留库", len(text_hold), "非卫衣/童装/套装等"],
        ["识图留库", len(vision_hold), "异色袖/袖印花/非黑白/不清等"],
        ["其它留库/详情失败", len(other_hold), ""],
        ["异常", len(errors), ""],
        ["模型", model, ""],
        ["实际识图条数", visioned, ""],
        ["识图张数 images_scanned", imgs, ""],
        ["cache_hits", cache, ""],
        ["prompt_tokens", pt, ""],
        ["completion_tokens", ct, ""],
        ["total_tokens", tt, ""],
        ["估费用 CNY", round(cny, 4), "粗算参考"],
        ["可用率", f"{(len(usable)/len(results)*100 if results else 0):.1f}%", "可用/已筛"],
    ]
    for r_i, row in enumerate(summary, 1):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if r_i == 1:
                cell.fill = header_fill
                cell.font = header_font
            elif c_i == 1:
                cell.fill = label_fill
    for i, w in enumerate([28, 22, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("留库原因")
    ws2.append(["原因", "条数"])
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
    for reason, n in reasons.most_common(40):
        ws2.append([reason, n])
    ws2.column_dimensions["A"].width = 50
    ws2.column_dimensions["B"].width = 10

    ws3 = wb.create_sheet("明细")
    headers = [
        "可用", "阶段", "原因", "货源ID", "commonId", "gmtCreate", "标题",
        "images", "prompt_tok", "completion_tok", "total_tok", "cny", "cache", "error",
    ]
    ws3.append(headers)
    for cell in ws3[1]:
        cell.fill = header_fill
        cell.font = header_font
    for r in results:
        ws3.append(
            [
                "是" if r.get("usable") else "否",
                r.get("stage"),
                r.get("reason"),
                r.get("source_id"),
                r.get("common_id"),
                r.get("gmt"),
                str(r.get("title") or "")[:120],
                r.get("images_scanned"),
                r.get("prompt_tokens"),
                r.get("completion_tokens"),
                r.get("total_tokens"),
                r.get("cny"),
                r.get("cache_hits"),
                r.get("error"),
            ]
        )
    for i, w in enumerate([6, 10, 28, 20, 14, 20, 36, 8, 12, 12, 12, 10, 8, 24], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    wb.save(path_prev)
    try:
        shutil.copy2(path_prev, path_desk)
    except Exception:
        pass
    return path_desk if path_desk.exists() else path_prev


def main() -> int:
    today0 = datetime.now(TZ).replace(hour=0, minute=0, second=0, microsecond=0)
    since = today0
    exclude_path: Path | None = None
    only_new_path: Path | None = None
    days = 1
    for a in sys.argv[1:]:
        if a.startswith("--exclude-baseline="):
            exclude_path = Path(a.split("=", 1)[1])
        elif a.startswith("--only-file="):
            only_new_path = Path(a.split("=", 1)[1])
        elif a.startswith("--days="):
            try:
                days = max(1, int(a.split("=", 1)[1]))
            except ValueError:
                days = 1
            since = today0 - timedelta(days=days - 1)
        elif a == "--all-success-today":
            since = today0

    settings = yaml.safe_load((ROOT / "config/settings.yaml").read_text(encoding="utf-8")) or {}
    m = settings.get("miaoshou") or {}
    paths = settings.get("paths") or {}
    out_dir = Path(paths.get("preview_dir") or ROOT / "data/previews")
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = out_dir / f"screen_hoodie_vision_{stamp}.log"
    result_path = out_dir / f"screen_hoodie_vision_{stamp}.json"

    def log(msg: str) -> None:
        line = f"{datetime.now():%Y-%m-%d %H:%M:%S} {msg}"
        print(line, flush=True)
        with log_path.open("a", encoding="utf-8") as f:
            f.write(line + "\n")

    if not SCHEME_PATH.exists():
        log(f"缺少 scheme: {SCHEME_PATH}")
        return 2

    client = MiaoshouClient(
        app_key=str(m.get("app_key") or ""),
        app_secret=str(m.get("app_secret") or ""),
        base_url=str(m.get("base_url") or ""),
        timeout=float(m.get("timeout_seconds") or 90),
        min_interval=1.0,
        max_retries=4,
    )
    scheme = json.loads(SCHEME_PATH.read_text(encoding="utf-8"))
    blank = Path(paths.get("blank_tee_dir") or ROOT / "assets/blank_tees")
    fill_urls = [str(u) for u in (settings.get("fill_image_urls") or []) if str(u).startswith("http")]
    vision_cfg = dict(settings.get("vision") or {})
    vision_cfg["enabled"] = True
    vision_cfg["product_type"] = "hoodie"
    vision_cfg["hold_on_tank_top"] = True
    vision_cfg["hold_on_contrast_sleeves"] = True
    model = str(vision_cfg.get("model") or "")

    if only_new_path and only_new_path.exists():
        payload = json.loads(only_new_path.read_text(encoding="utf-8"))
        pool = list(payload.get("rows") or [])
        log(f"从文件加载待筛 {len(pool)} 条 ← {only_new_path.name}")
    else:
        log(f"拉公共箱 success… since={since.isoformat()} days={days}")
        raw = list_success_since(client, since)
        log(f"success={len(raw)}，按标题预筛卫衣标记…")
        pool = [r for r in raw if title_looks_hoodie(str(r.get("title") or ""))]
        log(f"标题含卫衣标记 → {len(pool)}")
        if exclude_path and exclude_path.exists():
            base = json.loads(exclude_path.read_text(encoding="utf-8"))
            old = {str(x) for x in (base.get("source_ids") or [])}
            before = len(pool)
            pool = [r for r in pool if str(r.get("sid") or "") not in old]
            log(f"排除基线 {len(old)} 条后：{before} → {len(pool)}")

    if not pool:
        log("无卫衣候选，结束")
        write_report([], 0, stamp, out_dir, model)
        return 0

    import threading

    detail_lock = threading.Lock()
    results: list[dict] = []
    done_n = 0
    usable_n = 0
    t0 = time.time()

    def job(row: dict) -> dict:
        return screen_one(client, row, scheme, blank, fill_urls, vision_cfg, detail_lock)

    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = {ex.submit(job, r): r for r in pool}
        for fut in as_completed(futs):
            one = fut.result()
            results.append(one)
            done_n += 1
            if one.get("usable"):
                usable_n += 1
            if done_n % 10 == 0 or done_n == len(pool):
                pt = sum(int(x.get("prompt_tokens") or 0) for x in results)
                tt = sum(int(x.get("total_tokens") or 0) for x in results)
                cny = sum(float(x.get("cny") or 0) for x in results)
                log(
                    f"进度 {done_n}/{len(pool)} usable={usable_n} "
                    f"tok≈{tt or pt} cny≈{cny:.4f} "
                    f"elapsed={time.time()-t0:.0f}s"
                )
            if done_n % 20 == 0:
                result_path.write_text(
                    json.dumps(results, ensure_ascii=False, indent=2, default=str),
                    encoding="utf-8",
                )

    results.sort(key=lambda r: str(r.get("gmt") or ""), reverse=True)
    result_path.write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    usable_rows = [r for r in results if r.get("usable")]
    usable_queue = {
        "line": "hoodie",
        "scheme": "hoodie_mx",
        "stamp": stamp,
        "source_ids": [str(r.get("source_id")) for r in usable_rows],
        "rows": usable_rows,
        "count": len(usable_rows),
    }
    queue_path = out_dir / f"hoodie_usable_queue_{stamp}.json"
    queue_path.write_text(json.dumps(usable_queue, ensure_ascii=False, indent=2), encoding="utf-8")

    xlsx = write_report(results, len(pool), stamp, out_dir, model)
    usable = sum(1 for r in results if r.get("usable"))
    pt = sum(int(r.get("prompt_tokens") or 0) for r in results)
    ct = sum(int(r.get("completion_tokens") or 0) for r in results)
    tt = sum(int(r.get("total_tokens") or 0) for r in results) or (pt + ct)
    imgs = sum(int(r.get("images_scanned") or 0) for r in results)
    cny = sum(float(r.get("cny") or 0) for r in results)
    summary = {
        "line": "hoodie",
        "scheme": "hoodie_mx",
        "pool": len(pool),
        "screened": len(results),
        "usable": usable,
        "text_hold": sum(1 for r in results if r.get("stage") == "text"),
        "vision_hold": sum(1 for r in results if r.get("stage") == "vision"),
        "images_scanned": imgs,
        "prompt_tokens": pt,
        "completion_tokens": ct,
        "total_tokens": tt,
        "cny": round(cny, 4),
        "model": model,
        "xlsx": str(xlsx),
        "json": str(result_path),
        "queue": str(queue_path),
    }
    log(f"=== 结束 {json.dumps(summary, ensure_ascii=False)} ===")
    (out_dir / f"screen_hoodie_vision_{stamp}_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
