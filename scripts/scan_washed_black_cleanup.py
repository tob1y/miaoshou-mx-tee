# -*- coding: utf-8 -*-
"""两店已发布：昨天上架之前，按策略扫水洗黑/灰黑。

策略：
- 标题有黑，或选项名有明显黑/灰 → 核这些黑相关选项图，是否水洗/灰黑
- 标题无黑且选项名也无明显黑 → 只抽 1 张图；非灰黑则跳过
- gmtCreate < 2026-09-23
"""
from __future__ import annotations

import json
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from api.miaoshou_client import MiaoshouClient
from rules.vision_gate import (
    call_vision_body_qc,
    _qc_cache_get,
    _qc_cache_put,
)
from services.washed_black_feishu import WashedBlackFeishuSink

SHOPS = {18545217: "小韩1店", 18545044: "小赵1店"}
TZ = timezone(timedelta(hours=8))
# 昨天上架日起点：此前已上架的才扫
CUTOFF = datetime(2026, 9, 23, 0, 0, 0, tzinfo=TZ)
# 用户指定续跑起点（上次约 680）；与历史 max_scanned 取较大值
RESUME_FROM = 680
DONE_PATH = ROOT / "data/previews/washed_black_done_ids.json"

WASH_BAD = {"washed_black", "light_gray", "dark_gray"}
BLACK_FAMILY = {"pure_black", "washed_black", "light_gray", "dark_gray"}

# 文字无任何黑/灰相关 → 可跳过识图（省费用）；有黑相关才识图
_BLACKISH_TEXT = re.compile(
    r"(?i)negro|black|noir|schwarz|黑|lavad|washed|deslavad|gris|gray|grey|"
    r"charcoal|灰|炭黑|灰黑|水洗|graphite|antracit|anthracite"
)


def done_key(shop_id: int, detail_id: int) -> str:
    return f"{int(shop_id)}:{int(detail_id)}"


def load_done_ids(out_dir: Path, rows: list[dict]) -> set[str]:
    """已扫过的 shop:detail，避免重跑。

    来源：持久化文件 + 历史命中 JSON + 日志里出现过的 detail +
    历史最长一次进度（按同序列表前 N 条种子，补全未落盘的跳过项）。
    """
    done: set[str] = set()
    if DONE_PATH.exists():
        try:
            raw = json.loads(DONE_PATH.read_text(encoding="utf-8"))
            for x in raw.get("ids") or []:
                done.add(str(x))
        except Exception:
            pass

    max_scanned = 0
    for jp in out_dir.glob("washed_black_cleanup_*.json"):
        try:
            data = json.loads(jp.read_text(encoding="utf-8"))
        except Exception:
            continue
        for h in data.get("hits") or []:
            if h.get("detail_id") is None or h.get("shop_id") is None:
                continue
            done.add(done_key(int(h["shop_id"]), int(h["detail_id"])))
        sc = int((data.get("stats") or {}).get("scanned") or 0)
        if sc > max_scanned:
            max_scanned = sc

    for lp in out_dir.glob("washed_black_scan_*.log"):
        try:
            text = lp.read_text(encoding="utf-8", errors="replace")
        except Exception:
            continue
        # 日志多数只有 HIT/ERR 带 detail=；仍尽量收
        for m in re.finditer(r"detail=(\d+)", text):
            did = m.group(1)
            # 店未知时两店都记，宁多重跳不重扫
            done.add(done_key(18545044, int(did)))
            done.add(done_key(18545217, int(did)))

    # 同序列表前 N 条：覆盖最长一次已扫但未逐条落 id 的情况
    n = max(max_scanned, int(RESUME_FROM))
    if n > 0 and rows:
        n = min(n, len(rows))
        for r in rows[:n]:
            done.add(done_key(int(r["shop_id"]), int(r["detail_id"])))

    return done


def save_done_ids(done: set[str]) -> None:
    DONE_PATH.parent.mkdir(parents=True, exist_ok=True)
    DONE_PATH.write_text(
        json.dumps({"ids": sorted(done), "n": len(done)}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def parse_gmt(s: str) -> datetime | None:
    s = str(s or "").strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.strptime(s[:26], fmt).replace(tzinfo=TZ)
        except ValueError:
            continue
    return None


def text_has_blackish(product: dict, list_title: str) -> bool:
    blobs = [list_title, str(product.get("title") or "")]
    for prop in product.get("skuPropertyList") or []:
        if not isinstance(prop, dict):
            continue
        for v in prop.get("attrValueList") or []:
            if isinstance(v, dict):
                blobs.append(str(v.get("attrValue") or v.get("value") or ""))
    return bool(_BLACKISH_TEXT.search(" ".join(blobs)))


def list_published_before(client: MiaoshouClient, cutoff: datetime) -> list[dict]:
    """优先用发布成功记录（有 platformItemId + 上架时间）；不足时再补 published 箱。"""
    by_key: dict[tuple[int, int], dict] = {}

    for shop_id, shop_name in SHOPS.items():
        for page in range(1, 100):
            try:
                body = client.search_publish_tasks(page_no=page, page_size=100, shop_id=shop_id)
            except Exception:
                break
            data = body.get("data") or {}
            lst = data.get("moveCollectDetailList") or []
            if not lst:
                break
            page_kept = 0
            for it in lst:
                if not isinstance(it, dict):
                    continue
                st = str(it.get("status") or "").lower()
                if st and st not in ("success", "succeed", "ok"):
                    continue
                gmt = parse_gmt(str(it.get("gmtCreate") or it.get("gmtModified") or ""))
                if gmt is None or gmt >= cutoff:
                    continue
                did = it.get("collectBoxDetailId") or it.get("detailId")
                pid = it.get("platformItemId")
                if did is None or not pid:
                    continue
                key = (int(shop_id), int(did))
                by_key[key] = {
                    "detail_id": int(did),
                    "shop_id": int(shop_id),
                    "shop_name": shop_name,
                    "title": str(it.get("title") or "")[:200],
                    "source_id": str(it.get("sourceItemId") or it.get("itemNum") or ""),
                    "platform_item_id": str(pid),
                    "item_edit_url": str(it.get("itemEditUrl") or ""),
                    "gmt_create": str(it.get("gmtCreate") or ""),
                }
                page_kept += 1
            # 发布记录通常新→旧：整页都 >= cutoff（page_kept=0）且首页之后可继续翻，直到很旧
            oldest = parse_gmt(str(lst[-1].get("gmtCreate") or ""))
            if oldest and oldest < cutoff - timedelta(days=120):
                break
            if len(lst) < 100:
                break
            time.sleep(0.3)
    for page in range(1, 80):
        body = client.search_tiktok_box(page_no=page, page_size=100, status="published")
        batch = (body.get("data") or {}).get("detailList") or []
        if not batch:
            break
        for it in batch:
            shops = it.get("collectBoxDetailShopList") or []
            sids = [
                int(x["shopId"])
                for x in shops
                if isinstance(x, dict) and x.get("shopId") is not None
            ]
            our = [sid for sid in sids if sid in SHOPS]
            if not our:
                continue
            gmt = parse_gmt(str(it.get("gmtCreate") or it.get("gmtModified") or ""))
            if gmt is None or gmt >= cutoff:
                continue
            did = int(it["collectBoxDetailId"])
            key = (our[0], did)
            if key in by_key:
                continue
            by_key[key] = {
                "detail_id": did,
                "shop_id": our[0],
                "shop_name": SHOPS[our[0]],
                "title": str(it.get("title") or "")[:200],
                "source_id": str(it.get("itemNum") or ""),
                "platform_item_id": "",
                "item_edit_url": "",
                "gmt_create": str(it.get("gmtCreate") or ""),
            }
        if len(batch) < 100:
            break
        time.sleep(0.25)

    return sorted(by_key.values(), key=lambda r: (r["shop_id"], r["gmt_create"]))


def map_platform_item_ids(client: MiaoshouClient) -> dict[tuple[int, int], str]:
    """兼容旧调用；主列表已带 platform_item_id。"""
    return {}


def _option_image_rows(product: dict, max_n: int = 8) -> list[tuple[str, str, str]]:
    """返回 [(role, url, label), ...] 仅颜色轴选项图。"""
    rows: list[tuple[str, str, str]] = []
    props = product.get("skuPropertyList") or []
    color_prop = None
    for prop in props:
        if not isinstance(prop, dict):
            continue
        name = str(prop.get("name") or prop.get("attrName") or "").lower()
        if "talla" in name or "size" in name or "尺码" in name:
            continue
        if "color" in name or "colour" in name or "颜色" in name or name in (
            "color",
            "colores",
            "colour",
            "colours",
        ):
            color_prop = prop
            break
    if color_prop is None:
        for prop in props:
            if not isinstance(prop, dict):
                continue
            name = str(prop.get("name") or prop.get("attrName") or "").lower()
            if "talla" in name or "size" in name or "尺码" in name:
                continue
            color_prop = prop
            break
    if color_prop is None:
        return rows
    for i, v in enumerate(color_prop.get("attrValueList") or []):
        if not isinstance(v, dict):
            continue
        if str(v.get("isDelete") or "0") in ("1", "true", "True"):
            continue
        label = str(v.get("attrValue") or v.get("value") or "")
        img = ""
        for k in ("imgUrl", "imageUrl", "picUrl", "skuImage"):
            if v.get(k):
                img = str(v.get(k))
                break
        if not img:
            extras = v.get("supplementarySkuImageUrls")
            if isinstance(extras, list) and extras:
                img = str(extras[0] if not isinstance(extras[0], dict) else extras[0].get("url") or "")
        if not str(img).startswith("http"):
            continue
        rows.append((f"option_{i}", str(img), label))
        if len(rows) >= max_n:
            break
    return rows


def _title_blackish(title: str) -> bool:
    return bool(_BLACKISH_TEXT.search(title or ""))


def _label_blackish(label: str) -> bool:
    return bool(_BLACKISH_TEXT.search(label or ""))


def _first_any_image(product: dict) -> tuple[str, str, str] | None:
    """抽一张：优先第一张选项图，否则第一张主图。"""
    opts = _option_image_rows(product, max_n=1)
    if opts:
        return opts[0]
    for key in ("imgUrls", "main_images", "imageList", "images"):
        cur = product.get(key)
        if isinstance(cur, list) and cur:
            u = cur[0]
            if isinstance(u, dict):
                u = u.get("url") or u.get("imgUrl") or ""
            u = str(u or "")
            if u.startswith("http"):
                return ("main_0", u, "")
    thumb = str(product.get("thumbnail") or product.get("mainImage") or "")
    if thumb.startswith("http"):
        return ("main_0", thumb, "")
    return None


def _vision_call(url: str, vision_cfg: dict) -> dict:
    api_key = str(vision_cfg.get("api_key") or "")
    model = str(vision_cfg.get("model") or "gpt-5.6-luna")
    base_url = str(vision_cfg.get("base_url") or "https://api.openai.com/v1")
    timeout = float(vision_cfg.get("timeout_seconds") or 45)
    image_detail = vision_cfg.get("image_detail")
    if image_detail is None:
        image_detail = "low"
    use_cache = vision_cfg.get("cache_image_results", True) is not False
    vision = _qc_cache_get(url) if use_cache else None
    if vision is None:
        vision = call_vision_body_qc(
            url,
            api_key=api_key,
            model=model,
            base_url=base_url,
            timeout=timeout,
            image_detail=image_detail,
        )
        if use_cache and "error" not in vision:
            _qc_cache_put(url, vision)
    return vision


def scan_one(
    product: dict,
    *,
    vision_cfg: dict,
    list_title: str = "",
    max_images: int = 8,
) -> dict:
    """两种模式：

    A) 标题有黑，或选项名有明显黑/灰 → 扫这些黑相关选项图，查水洗/灰黑
    B) 标题无黑且选项名也无明显黑 → 只抽 1 张图；非灰黑则跳过
    """
    title = str(product.get("title") or list_title or "")
    title_blk = _title_blackish(title)
    max_opt = int(vision_cfg.get("max_option_images") or max_images or 8)
    opt_rows = _option_image_rows(product, max_n=max_opt)

    blackish_opts = [(r, u, lab) for r, u, lab in opt_rows if _label_blackish(lab)]
    # 无色名选项：标题带黑时一并核（可能是款号当色名）
    unlabeled = [(r, u, lab) for r, u, lab in opt_rows if not (lab or "").strip()]
    mode_full = title_blk or bool(blackish_opts)

    if mode_full:
        to_scan = list(blackish_opts)
        if title_blk and unlabeled:
            to_scan.extend(unlabeled)
        # 标题有黑但选项名都没写成黑：仍核全部有图选项（防漏水洗黑）
        if title_blk and not blackish_opts:
            to_scan = list(opt_rows) if opt_rows else []
        # 去重保序
        seen: set[str] = set()
        ordered: list[tuple[str, str, str]] = []
        for item in to_scan:
            if item[1] in seen:
                continue
            seen.add(item[1])
            ordered.append(item)
        mode = "full_black_options"
    else:
        sample = _first_any_image(product)
        ordered = [sample] if sample else []
        mode = "sample_one"

    if not ordered:
        return {
            "hit": False,
            "skip_no_black": True,
            "reasons": [],
            "hits": [],
            "images_scanned": 0,
            "saw_black": False,
            "early_stop": False,
            "mode": mode,
        }

    saw_black = False
    hits: list[dict] = []
    scanned = 0

    for role, url, label in ordered:
        vision = _vision_call(url, vision_cfg)
        scanned += 1
        bc = str(vision.get("body_color") or "")
        entry = {
            "role": role,
            "label": label,
            "body_color": bc,
            "confident": bool(vision.get("confident")),
            "reason": str(vision.get("reason") or "")[:160],
            "url": url[:180],
            "cache_hit": bool(vision.get("cache_hit")),
            "mode": mode,
        }
        if bc in BLACK_FAMILY:
            saw_black = True
        if bc in WASH_BAD:
            hits.append({**entry, "flag": bc})
            return {
                "hit": True,
                "skip_no_black": False,
                "reasons": [bc],
                "hits": hits,
                "images_scanned": scanned,
                "saw_black": True,
                "early_stop": True,
                "mode": mode,
            }

    # sample 模式：抽一张不是灰黑 → 跳过
    if mode == "sample_one":
        return {
            "hit": False,
            "skip_no_black": True,
            "reasons": [],
            "hits": [],
            "images_scanned": scanned,
            "saw_black": saw_black,
            "early_stop": False,
            "mode": mode,
        }

    if not saw_black:
        return {
            "hit": False,
            "skip_no_black": True,
            "reasons": [],
            "hits": [],
            "images_scanned": scanned,
            "saw_black": False,
            "early_stop": False,
            "mode": mode,
        }

    return {
        "hit": False,
        "skip_no_black": False,
        "reasons": [],
        "hits": [],
        "images_scanned": scanned,
        "saw_black": True,
        "early_stop": False,
        "mode": mode,
    }


def _safe_copy_xlsx(src: Path, *dests: Path) -> None:
    """桌面文件被 Excel 占用时写失败则改写「_更新中」副本。"""
    import shutil

    for dest in dests:
        tmp = dest.with_suffix(dest.suffix + ".tmp")
        try:
            shutil.copy2(src, tmp)
            tmp.replace(dest)
        except Exception:
            try:
                if tmp.exists():
                    tmp.unlink(missing_ok=True)
            except Exception:
                pass
            try:
                alt = dest.with_name(dest.stem + "_更新中" + dest.suffix)
                shutil.copy2(src, alt)
            except Exception:
                pass


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "水洗黑清理清单"
    headers = [
        "序号",
        "标题",
        "店铺",
        "shop_id",
        "商品ID(platformItemId)",
        "detailId",
        "货源ID",
        "gmtCreate",
        "命中原因",
        "扫图张数",
        "命中说明",
        "编辑链接",
    ]
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, horizontal="center")
        cell.border = thin
    for i, r in enumerate(rows, 1):
        hit_note = "; ".join(
            f"{h.get('role')}:{h.get('flag') or h.get('body_color')}"
            for h in (r.get("hits") or [])[:6]
        )
        vals = [
            i,
            r.get("title") or "",
            r.get("shop_name") or "",
            r.get("shop_id"),
            r.get("platform_item_id") or "",
            r.get("detail_id"),
            r.get("source_id") or "",
            r.get("gmt_create") or "",
            ",".join(r.get("reasons") or []),
            r.get("images_scanned"),
            hit_note,
            r.get("item_edit_url") or "",
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(i + 1, col, v)
            cell.border = thin
    from openpyxl.utils import get_column_letter

    widths = [6, 40, 12, 12, 22, 14, 22, 20, 28, 10, 40, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("说明")
    ws2["A1"] = "口径"
    ws2["A2"] = "范围：两店 published 且 gmtCreate < 2026-09-23（昨天上架前）"
    ws2["A3"] = "标题有黑或选项名有黑/灰：核黑相关选项图，命中 washed/light_gray/dark_gray 则记入"
    ws2["A4"] = "标题无黑且选项名也无明显黑：只抽1张图，非灰黑则跳过"
    ws2["A5"] = "异色袖不记"
    ws2["A6"] = "商品ID：platformItemId；缺则 detailId/货源ID 补查"
    wb.save(path)


def main() -> int:
    limit = 0  # 0=全部
    only: str | None = None
    global RESUME_FROM
    for a in sys.argv[1:]:
        if a.startswith("--limit="):
            limit = int(a.split("=", 1)[1])
        elif a.startswith("--only="):
            only = a.split("=", 1)[1].strip()
        elif a.startswith("--from="):
            RESUME_FROM = max(0, int(a.split("=", 1)[1]))

    settings = yaml.safe_load((ROOT / "config/settings.yaml").read_text(encoding="utf-8")) or {}
    m = settings.get("miaoshou") or {}
    fcfg = settings.get("feishu") or {}
    vision_cfg = dict(settings.get("vision") or {})
    out_dir = Path((settings.get("paths") or {}).get("preview_dir") or ROOT / "data/previews")
    out_dir.mkdir(parents=True, exist_ok=True)
    desk = Path.home() / "Desktop"
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = out_dir / f"washed_black_scan_{stamp}.log"

    def log(msg: str) -> None:
        line = f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')} {msg}"
        print(line, flush=True)
        with log_path.open("a", encoding="utf-8") as f:
            f.write(line + "\n")

    feishu_sink: WashedBlackFeishuSink | None = None
    wb_token = str(fcfg.get("washed_black_sheet_token") or "").strip()
    if wb_token and fcfg.get("app_id") and fcfg.get("app_secret"):
        feishu_sink = WashedBlackFeishuSink(
            str(fcfg["app_id"]),
            str(fcfg["app_secret"]),
            wb_token,
            log=log,
        )
        if feishu_sink.ensure():
            log("飞书水洗黑表已接通，命中将实时追加")
        else:
            log("飞书水洗黑表暂不可用，命中仅本地/桌面落盘")

    client = MiaoshouClient(
        app_key=str(m.get("app_key") or ""),
        app_secret=str(m.get("app_secret") or ""),
        base_url=str(m.get("base_url") or ""),
        timeout=float(m.get("timeout_seconds") or 90),
        min_interval=1.0,
        max_retries=4,
    )

    log(f"拉 published（gmtCreate < {CUTOFF.date()}）…")
    rows = list_published_before(client, CUTOFF)
    if only == "han":
        rows = [r for r in rows if r["shop_id"] == 18545217]
    elif only == "zhao":
        rows = [r for r in rows if r["shop_id"] == 18545044]
    if limit > 0:
        rows = rows[:limit]

    by_shop: dict[str, int] = {}
    for r in rows:
        by_shop[r["shop_name"]] = by_shop.get(r["shop_name"], 0) + 1
    log(f"待扫池 {len(rows)} {by_shop}")

    done = load_done_ids(out_dir, rows)
    before = len(rows)
    rows = [r for r in rows if done_key(int(r["shop_id"]), int(r["detail_id"])) not in done]
    log(f"已扫跳过 {before - len(rows)}（done={len(done)}）剩余 {len(rows)}")
    save_done_ids(done)

    # 续跑：合并历史命中，桌面表不丢；飞书只追加本轮新命中
    hist_hits: list[dict] = []
    seen_hit: set[str] = set()
    for jp in sorted(out_dir.glob("washed_black_cleanup_*.json"), key=lambda p: p.stat().st_mtime):
        try:
            data = json.loads(jp.read_text(encoding="utf-8"))
        except Exception:
            continue
        for h in data.get("hits") or []:
            k = done_key(int(h.get("shop_id") or 0), int(h.get("detail_id") or 0))
            if k in seen_hit:
                continue
            seen_hit.add(k)
            hist_hits.append(h)

    hits: list[dict] = list(hist_hits)
    stats = {
        "scanned": 0,
        "hit": 0,
        "skip_sample_ok": 0,
        "skip_no_black_vision": 0,
        "skip_pure_ok": 0,
        "mode_full": 0,
        "mode_sample": 0,
        "err": 0,
        "deleted": 0,
        "skipped_done": before - len(rows),
        "hist_hits": len(hist_hits),
    }

    json_path = out_dir / f"washed_black_cleanup_{stamp}.json"
    xlsx_name = f"水洗黑清理清单_{stamp}.xlsx"
    xlsx_path = out_dir / xlsx_name

    for i, row in enumerate(rows, 1):
        stats["scanned"] += 1
        dk = done_key(int(row["shop_id"]), int(row["detail_id"]))
        if i % 10 == 0 or i == 1:
            log(
                f"进度 {i}/{len(rows)} hit={stats['hit']} "
                f"full={stats['mode_full']} sample={stats['mode_sample']} "
                f"ok_pure={stats['skip_pure_ok']} sample_ok={stats['skip_sample_ok']} err={stats['err']}"
            )
            # 增量落盘（项目目录 + 桌面同名，方便边扫边看）
            json_path.write_text(
                json.dumps({"stats": stats, "hits": hits}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if hits:
                try:
                    write_xlsx(hits, xlsx_path)
                    _safe_copy_xlsx(xlsx_path, desk / "水洗黑清理清单.xlsx")
                except Exception as e:
                    log(f"桌面增量拷贝失败: {e}")

        try:
            try:
                client.claim_to_shops([row["detail_id"]], [row["shop_id"]])
            except Exception:
                pass
            time.sleep(0.15)
            _oss, product = client.get_shop_detail(row["detail_id"], row["shop_id"])
        except Exception as e:
            if "已被删除" in str(e):
                stats["deleted"] += 1
                done.add(dk)
                if i % 10 == 0:
                    save_done_ids(done)
                continue
            stats["err"] += 1
            log(f"ERR detail={row['detail_id']}: {e}")
            done.add(dk)
            if i % 10 == 0:
                save_done_ids(done)
            continue

        title = str(product.get("title") or row["title"] or "")[:200]
        # 缺商品ID时：用货源ID再查发布记录补 platformItemId
        if not row.get("platform_item_id") and row.get("source_id"):
            try:
                sid = str(row["source_id"])
                body = client.search_publish_tasks(
                    page_no=1,
                    page_size=20,
                    shop_id=int(row["shop_id"]),
                    extra_filter={"sourceItemIdKeyword": sid},
                )
                for it in (body.get("data") or {}).get("moveCollectDetailList") or []:
                    src = str(it.get("sourceItemId") or it.get("itemNum") or "")
                    pid = str(it.get("platformItemId") or "")
                    if pid and src == sid:
                        row["platform_item_id"] = pid
                        row["item_edit_url"] = str(it.get("itemEditUrl") or "")
                        break
                    did = it.get("collectBoxDetailId") or it.get("detailId")
                    if pid and did is not None and int(did) == int(row["detail_id"]):
                        row["platform_item_id"] = pid
                        row["item_edit_url"] = str(it.get("itemEditUrl") or "")
                        break
            except Exception:
                pass

        try:
            result = scan_one(product, vision_cfg=vision_cfg, list_title=title)
        except Exception as e:
            stats["err"] += 1
            log(f"VISION_ERR detail={row['detail_id']}: {e}")
            done.add(dk)
            if i % 10 == 0:
                save_done_ids(done)
            continue

        if result.get("mode") == "sample_one":
            stats["mode_sample"] += 1
        else:
            stats["mode_full"] += 1

        if result.get("skip_no_black"):
            if result.get("mode") == "sample_one":
                stats["skip_sample_ok"] += 1
            else:
                stats["skip_no_black_vision"] += 1
            done.add(dk)
            if i % 10 == 0:
                save_done_ids(done)
            continue
        if not result.get("hit"):
            stats["skip_pure_ok"] += 1
            done.add(dk)
            if i % 10 == 0:
                save_done_ids(done)
            continue

        stats["hit"] += 1
        rec = {
            **row,
            "title": title,
            "platform_item_id": row.get("platform_item_id") or "",
            "item_edit_url": row.get("item_edit_url") or "",
            "reasons": result.get("reasons") or [],
            "hits": result.get("hits") or [],
            "images_scanned": result.get("images_scanned"),
            "early_stop": result.get("early_stop"),
            "mode": result.get("mode"),
        }
        hits.append(rec)
        log(
            f"HIT {row['shop_name']} detail={row['detail_id']} "
            f"product={rec['platform_item_id'] or '-'} reasons={rec['reasons']} | {title[:60]}"
        )
        if feishu_sink is not None:
            n = feishu_sink.append_hits([rec])
            if n:
                log(f"已同步飞书 +{n}")
        # 每命中一条立刻刷表到桌面
        try:
            json_path.write_text(
                json.dumps({"stats": stats, "hits": hits}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            write_xlsx(hits, xlsx_path)
            _safe_copy_xlsx(xlsx_path, desk / "水洗黑清理清单.xlsx")
        except Exception as e:
            log(f"HIT落盘失败: {e}")
        done.add(dk)
        if i % 10 == 0:
            save_done_ids(done)

    json_path.write_text(
        json.dumps(
            {
                "cutoff": str(CUTOFF),
                "stats": stats,
                "hits": hits,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    write_xlsx(hits, xlsx_path)
    desk_xlsx = desk / "水洗黑清理清单.xlsx"
    try:
        _safe_copy_xlsx(xlsx_path, desk_xlsx)
    except Exception as e:
        log(f"桌面拷贝失败: {e}")
        desk_xlsx = xlsx_path

    log(
        f"=== DONE scanned={stats['scanned']} hit={stats['hit']} "
        f"full={stats['mode_full']} sample={stats['mode_sample']} "
        f"ok_pure={stats['skip_pure_ok']} sample_ok={stats['skip_sample_ok']} err={stats['err']} "
        f"skipped_done={stats.get('skipped_done')} done_total={len(done)} "
        f"→ {xlsx_path.name} / Desktop"
    )
    save_done_ids(done)
    print("XLSX", desk_xlsx)
    print("JSON", json_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
