# -*- coding: utf-8 -*-
"""今日上架 + 水洗黑扫描汇总报表 → 桌面 Excel。"""
from __future__ import annotations

import json
import re
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
PREV = ROOT / "data/previews"
DESK = Path.home() / "Desktop"
QUOTA = ROOT / "data/daily_quota.json"
DONE_PATH = PREV / "washed_black_done_ids.json"
ZHAO, HAN = "18545044", "18545217"
TODAY = datetime.now().strftime("%Y-%m-%d")
TODAY_COMPACT = datetime.now().strftime("%Y%m%d")


def load_quota() -> dict:
    try:
        return (json.loads(QUOTA.read_text(encoding="utf-8")).get(TODAY) or {})
    except Exception:
        return {}


def count_log_patterns(path: Path) -> dict[str, int]:
    if not path.exists():
        return {}
    t = path.read_text(encoding="utf-8", errors="replace")
    return {
        "ok": len(re.findall(r"OK 已上架", t)),
        "text_hold": len(re.findall(r"文字留库", t)),
        "vision_hold": len(re.findall(r"识图/改品留库|留库不上架", t)),
        "fail": len(re.findall(r"\bFAIL\b", t)),
        "unclear": len(re.findall(r"main_vision_unclear", t)),
    }


def latest_logs(glob_pat: str) -> list[Path]:
    return sorted(PREV.glob(glob_pat), key=lambda p: p.stat().st_mtime)


def wash_stats() -> dict:
    hit_map: dict[str, dict] = {}
    max_scanned = 0
    runs = []
    for jp in sorted(PREV.glob("washed_black_cleanup_*.json"), key=lambda p: p.stat().st_mtime):
        try:
            data = json.loads(jp.read_text(encoding="utf-8"))
        except Exception:
            continue
        st = data.get("stats") or {}
        sc = int(st.get("scanned") or 0)
        max_scanned = max(max_scanned, sc)
        runs.append(
            {
                "file": jp.name,
                "scanned": sc,
                "hits": len(data.get("hits") or []),
                "skipped_done": st.get("skipped_done"),
                "vision_calls": st.get("vision_calls") or st.get("vision_n"),
                "cny": st.get("cny") or st.get("vision_cny"),
            }
        )
        for h in data.get("hits") or []:
            k = f"{h.get('shop_id')}:{h.get('detail_id')}"
            hit_map[k] = h

    done_n = 0
    if DONE_PATH.exists():
        try:
            done_n = int((json.loads(DONE_PATH.read_text(encoding="utf-8")) or {}).get("n") or 0)
        except Exception:
            pass

    hits = list(hit_map.values())
    by_shop = Counter(str(h.get("shop_name") or "") for h in hits)
    by_reason = Counter()
    for h in hits:
        rs = h.get("reasons") or []
        if isinstance(rs, list) and rs:
            for r in rs:
                by_reason[str(r)] += 1
        else:
            by_reason["unknown"] += 1

    # 从最新日志估用量
    vision_calls = 0
    cny = 0.0
    for lp in PREV.glob("washed_black_scan_*.log"):
        t = lp.read_text(encoding="utf-8", errors="replace")
        for m in re.finditer(r"vision[_\s]*calls?=(\d+)", t, re.I):
            vision_calls = max(vision_calls, int(m.group(1)))
        for m in re.finditer(r"cny[=: ]([0-9.]+)", t, re.I):
            try:
                cny = max(cny, float(m.group(1)))
            except ValueError:
                pass
    for r in runs:
        if r.get("vision_calls"):
            try:
                vision_calls = max(vision_calls, int(r["vision_calls"]))
            except Exception:
                pass
        if r.get("cny"):
            try:
                cny = max(cny, float(r["cny"]))
            except Exception:
                pass

    # 汇总 stats 字段（若有）
    total_vision = 0
    total_cny = 0.0
    for jp in PREV.glob("washed_black_cleanup_*.json"):
        try:
            st = (json.loads(jp.read_text(encoding="utf-8")).get("stats") or {})
        except Exception:
            continue
        for k in ("vision_calls", "vision_n", "images_scanned"):
            if st.get(k) is not None:
                try:
                    total_vision += int(st[k])
                except Exception:
                    pass
        for k in ("cny", "vision_cny", "cost_cny"):
            if st.get(k) is not None:
                try:
                    total_cny += float(st[k])
                except Exception:
                    pass

    return {
        "done_n": done_n,
        "scanned_est": done_n or max_scanned,
        "hits": hits,
        "hits_n": len(hits),
        "by_shop": dict(by_shop),
        "by_reason": dict(by_reason),
        "runs": runs,
        "vision_calls_est": total_vision or vision_calls,
        "cny_est": round(total_cny or cny, 4),
    }


def style_header(ws, row: int, cols: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def main() -> int:
    q = load_quota()
    zhao_n = int(q.get(ZHAO) or 0)
    han_n = int(q.get(HAN) or 0)

    # 今日并发日志汇总
    zhao_pat = count_log_patterns(PREV / "zhao_dual_stdout.log")
    han_pat = count_log_patterns(PREV / "han_dual_stdout.log")
    for lp in latest_logs(f"zhao_concurrent_{TODAY_COMPACT}_*.log")[-3:]:
        for k, v in count_log_patterns(lp).items():
            zhao_pat[k] = zhao_pat.get(k, 0) + v
    for lp in latest_logs(f"han_concurrent_{TODAY_COMPACT}_*.log")[-3:]:
        for k, v in count_log_patterns(lp).items():
            han_pat[k] = han_pat.get(k, 0) + v

    wash = wash_stats()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_name = f"今日上架与水洗黑报告_{TODAY_COMPACT}.xlsx"
    out_prev = PREV / out_name
    out_desk = DESK / out_name

    wb = Workbook()
    ws = wb.active
    ws.title = "今日上架"
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    label_fill = PatternFill("solid", fgColor="D6EAF8")

    rows_pub = [
        ["口径", "小赵1店", "小韩1店", "合计", "说明"],
        ["妙手日配额（今日 success）", zhao_n, han_n, zhao_n + han_n, "daily_quota.json"],
        ["日志 OK 已上架（参考）", zhao_pat.get("ok", 0), han_pat.get("ok", 0), zhao_pat.get("ok", 0) + han_pat.get("ok", 0), "可能含多轮日志累加"],
        ["文字留库", zhao_pat.get("text_hold", 0), han_pat.get("text_hold", 0), zhao_pat.get("text_hold", 0) + han_pat.get("text_hold", 0), ""],
        ["识图/改品留库", zhao_pat.get("vision_hold", 0), han_pat.get("vision_hold", 0), zhao_pat.get("vision_hold", 0) + han_pat.get("vision_hold", 0), "含 unclear"],
        ["main_vision_unclear", zhao_pat.get("unclear", 0), han_pat.get("unclear", 0), zhao_pat.get("unclear", 0) + han_pat.get("unclear", 0), ""],
        ["FAIL", zhao_pat.get("fail", 0), han_pat.get("fail", 0), zhao_pat.get("fail", 0) + han_pat.get("fail", 0), ""],
        ["单店日限", 300, 300, 600, ""],
    ]
    for r_i, row in enumerate(rows_pub, 1):
        for c_i, v in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, v)
            cell.border = thin
            if r_i == 1:
                pass
            elif c_i == 1:
                cell.fill = label_fill
    style_header(ws, 1, 5)
    for i, w in enumerate([28, 14, 14, 12, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("水洗黑扫描")
    wash_rows = [
        ["指标", "数值", "说明"],
        ["已扫（done_ids）", wash["scanned_est"], "持久化已扫 shop:detail，约等于累计扫描数"],
        ["去重命中数", wash["hits_n"], "历史 cleanup JSON 合并"],
        ["识图调用估", wash["vision_calls_est"], "各轮 stats/日志累加或峰值"],
        ["费用估 CNY", wash["cny_est"], "若日志有 cny"],
        ["本轮策略", "只扫小赵；跳过小韩", "小韩店内已自检"],
        ["命中→飞书", "实时追加 washed_black 表", "sheet token 配置"],
    ]
    for r_i, row in enumerate(wash_rows, 1):
        for c_i, v in enumerate(row, 1):
            cell = ws2.cell(r_i, c_i, v)
            cell.border = thin
            if c_i == 1 and r_i > 1:
                cell.fill = label_fill
    style_header(ws2, 1, 3)
    r = len(wash_rows) + 2
    ws2.cell(r, 1, "按店铺命中")
    ws2.cell(r, 1).font = Font(bold=True)
    r += 1
    ws2.cell(r, 1, "店铺")
    ws2.cell(r, 2, "命中数")
    style_header(ws2, r, 2)
    r += 1
    for shop, n in sorted(wash["by_shop"].items(), key=lambda x: -x[1]):
        ws2.cell(r, 1, shop)
        ws2.cell(r, 2, n)
        r += 1
    r += 1
    ws2.cell(r, 1, "按原因命中")
    ws2.cell(r, 1).font = Font(bold=True)
    r += 1
    ws2.cell(r, 1, "原因")
    ws2.cell(r, 2, "条数")
    style_header(ws2, r, 2)
    r += 1
    for reason, n in sorted(wash["by_reason"].items(), key=lambda x: -x[1]):
        ws2.cell(r, 1, reason)
        ws2.cell(r, 2, n)
        r += 1
    for i, w in enumerate([28, 16, 40], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws3 = wb.create_sheet("水洗黑命中明细")
    headers = ["店铺", "商品ID", "detailId", "货源ID", "命中原因", "gmtCreate", "标题", "编辑链接"]
    for c, h in enumerate(headers, 1):
        ws3.cell(1, c, h)
    style_header(ws3, 1, len(headers))
    for r_i, h in enumerate(
        sorted(wash["hits"], key=lambda x: (str(x.get("shop_name")), str(x.get("gmt_create")))),
        2,
    ):
        ws3.cell(r_i, 1, h.get("shop_name") or "")
        ws3.cell(r_i, 2, h.get("platform_item_id") or "")
        ws3.cell(r_i, 3, h.get("detail_id") or "")
        ws3.cell(r_i, 4, h.get("source_id") or "")
        rs = h.get("reasons") or []
        ws3.cell(r_i, 5, ",".join(rs) if isinstance(rs, list) else str(rs))
        ws3.cell(r_i, 6, h.get("gmt_create") or "")
        ws3.cell(r_i, 7, str(h.get("title") or "")[:120])
        ws3.cell(r_i, 8, h.get("item_edit_url") or "")
    for i, w in enumerate([12, 22, 14, 22, 28, 20, 40, 36], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws4 = wb.create_sheet("说明")
    ws4["A1"] = f"生成时间 {datetime.now():%Y-%m-%d %H:%M:%S}"
    ws4["A2"] = "上架：以 daily_quota 为权威件数；日志留库为参考。"
    ws4["A3"] = "水洗黑：此前两店已扫约 done_ids 条；后续仅补扫小赵未扫项，命中追加飞书表。"
    ws4["A4"] = "异色袖不作为独立命中（历史若有 contrast_sleeves 会留在明细里）。"
    ws4["A5"] = f"meta stamp={stamp} done={wash['scanned_est']} hits={wash['hits_n']}"

    wb.save(out_prev)
    try:
        shutil.copy2(out_prev, out_desk)
        desk_ok = str(out_desk)
    except Exception as e:
        desk_ok = f"桌面复制失败: {e}"
    meta = {
        "stamp": stamp,
        "quota": {"zhao": zhao_n, "han": han_n},
        "wash": {
            "scanned_est": wash["scanned_est"],
            "hits_n": wash["hits_n"],
            "by_shop": wash["by_shop"],
            "by_reason": wash["by_reason"],
            "vision_calls_est": wash["vision_calls_est"],
            "cny_est": wash["cny_est"],
        },
        "xlsx_prev": str(out_prev),
        "xlsx_desk": desk_ok,
    }
    (PREV / f"today_publish_wash_report_meta_{stamp}.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(meta, ensure_ascii=False, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
